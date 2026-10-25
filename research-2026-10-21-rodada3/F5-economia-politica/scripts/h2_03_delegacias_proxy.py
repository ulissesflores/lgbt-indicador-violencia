"""H2 sensibilidade: o ESTADIC traz o TOTAL de delegacias especializadas da Polícia Civil por UF (ESEG5134 em 2019; sem tipo: inclui DEAM etc.), não delegacia de crime racial/LGBT.
Proxy fraco, reportado como tal. Teste: Spearman entre nº de delegacias especializadas por milhão de habitantes e nº de UF-anos sem informação no Anuário; permutação 20000, semente fixa.
Rodar: /opt/homebrew/bin/python3.14 -I scripts/h2_03_delegacias_proxy.py (depois de h2_02_analise.py)"""
import os as _os
import pathlib as _pl

REPO_ROOT = _os.environ.get("REPO_ROOT") or str(_pl.Path(__file__).resolve().parents[3])  # raiz configurável
import openpyxl, numpy as np, pandas as pd
from scipy import stats
B = (REPO_ROOT + "/research-2026-10-21-rodada3/F5-economia-politica")
out = {}
for ed, f in ((2019, "estadic2019.xlsx"), (2023, "estadic2023.xlsx")):
    wb = openpyxl.load_workbook(f"{B}/dados/raw/estadic/{f}", read_only=True, data_only=True)
    ws = wb["Segurança pública"] if "Segurança pública" in wb.sheetnames else wb["Segurança Pública"]
    rows = list(ws.iter_rows(values_only=True)); h = [str(c) for c in rows[0]]
    if "ESEG5134" not in h: print(ed, "sem ESEG5134"); continue
    j = h.index("ESEG5134"); out[ed] = {int(r[0]): r[j] for r in rows[1:]}
U = pd.read_csv(f"{B}/dados/h2_uf_resumo.csv", index_col=0)
at = pd.read_csv((REPO_ROOT + "/dados/atlas_uf_series.csv"))
t = at[at.variavel == "homicidios_total_taxa"].merge(at[at.variavel == "homicidios_total_n"], on=["codigo_uf", "ano"], suffixes=("_t", "_n")); t["pop"] = t.valor_n / t.valor_t * 1e5
pop = t[t.ano == 2022].set_index("codigo_uf")["pop"]
ed = 2019; v = pd.Series(out[ed]).apply(pd.to_numeric, errors="coerce"); U["delegesp"] = v.reindex(U.index); U["delegesp_pmh"] = U.delegesp / pop.reindex(U.index) * 1e6
ok = U.dropna(subset=["delegesp_pmh"]); r = stats.spearmanr(ok.Z, ok.delegesp_pmh)
rng = np.random.default_rng(20261026); pp = np.array([stats.spearmanr(rng.permutation(ok.Z.values), ok.delegesp_pmh)[0] for _ in range(20000)])
print("n UFs com valor:", len(ok), "| Spearman Z x delegacias especializadas/milhao hab (ESTADIC 2019):", round(r[0], 3), "| p perm:", round(float(np.mean(np.abs(pp) >= abs(r[0]))), 3))
pd.DataFrame([dict(teste="Spearman Z x delegacias especializadas por milhão de hab (ESEG5134, ESTADIC 2019; sem tipo)", rho=r[0], p_perm=float(np.mean(np.abs(pp) >= abs(r[0]))), n=len(ok))]).to_csv(f"{B}/dados/h2_resultados_delegacias_proxy.csv", index=False)
