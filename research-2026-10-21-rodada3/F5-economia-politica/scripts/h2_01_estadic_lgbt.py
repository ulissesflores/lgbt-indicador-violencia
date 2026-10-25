"""H2 insumo: estrutura estadual para enfrentamento da LGBTfobia, ESTADIC/IBGE 2019 e 2023 (bases oficiais .xlsx do FTP do IBGE; sha256 em dados/raw/estadic/SHA256SUMS.txt).
Variáveis (códigos IBGE do dicionário das próprias bases): EDHU086 (órgão de direitos humanos executa programas/ações para LGBT), EDHU5410 (comitê gestor de enfrentamento à violência LGBTfóbica),
EDHU576 (política/programa de promoção de direitos LGBT e/ou enfrentamento à LGBTfobia), ESEG4413 (PM: capacitação com tema LGBTfobia), ESEG487 (ações/projetos de prevenção com tema LGBTfobia), ESEG5913 (Polícia Civil: capacitação com tema LGBTfobia).
Saída: dados/h2_estadic_lgbt_uf.csv (uma linha por UF x edição, valores brutos 'Sim'/'Não'/código do IBGE).
Rodar: /opt/homebrew/bin/python3.14 -I scripts/h2_01_estadic_lgbt.py"""
import os as _os
import pathlib as _pl

REPO_ROOT = _os.environ.get("REPO_ROOT") or str(_pl.Path(__file__).resolve().parents[3])  # raiz configurável
import openpyxl, pandas as pd
B = (REPO_ROOT + "/research-2026-10-21-rodada3/F5-economia-politica")
VARS = {"EDHU086": "Direitos humanos|Direitos Humanos", "EDHU5410": "Direitos humanos|Direitos Humanos", "EDHU576": "Direitos humanos|Direitos Humanos",
        "ESEG4413": "Segurança pública|Segurança Pública", "ESEG487": "Segurança pública|Segurança Pública", "ESEG5913": "Segurança pública|Segurança Pública"}
out = []
for ed, f in ((2019, "estadic2019.xlsx"), (2023, "estadic2023.xlsx")):
    wb = openpyxl.load_workbook(f"{B}/dados/raw/estadic/{f}", read_only=True, data_only=True)
    tab = {}
    for v, sheets in VARS.items():
        nome = next(s for s in sheets.split("|") if s in wb.sheetnames)
        rows = list(wb[nome].iter_rows(values_only=True)); hdr = [str(c) for c in rows[0]]
        j = hdr.index(v); cu = 0
        for r in rows[1:]:
            tab.setdefault(int(r[0]), {})[v] = r[j]
    for cod, d in tab.items(): out.append(dict(edicao=ed, cod_uf=cod, **d))
df = pd.DataFrame(out); df.to_csv(f"{B}/dados/h2_estadic_lgbt_uf.csv", index=False)
print(df.groupby("edicao").agg(lambda s: s.value_counts(dropna=False).to_dict()).T.to_string())
